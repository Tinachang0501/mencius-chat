# -*- coding: utf-8 -*-
"""
孟子 AI 古人聊天室｜A4 雙面自學學習單產生器
產出：output/word/孟子學習單_學生版.docx、孟子學習單_教師解答版.docx
PDF 轉檔請執行 scripts/export_pdf.ps1（使用本機 Microsoft Word）。
"""
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "output" / "word"
AVATAR = ROOT / "assets" / "avatar.jpg"

BODY_FONT = "標楷體"
HEAD_FONT = "微軟正黑體"
NAVY = RGBColor(0x2E, 0x3F, 0x5C)
RED = RGBColor(0xC0, 0x1F, 0x1F)   # 朱紅
GRAY = RGBColor(0x55, 0x55, 0x55)
BODY_PT = 11


# ---------------------------------------------------------------- 基本工具
def set_font(run, size=BODY_PT, bold=False, color=None, font=BODY_FONT, underline=False):
    run.font.name = font
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.underline = underline
    if color is not None:
        run.font.color.rgb = color
    return run


def para(doc_or_cell, align=None, after=2, before=0, line=1.2):
    p = doc_or_cell.add_paragraph()
    pf = p.paragraph_format
    pf.space_after = Pt(after)
    pf.space_before = Pt(before)
    pf.line_spacing = line
    if align:
        p.alignment = align
    return p


def shade(cell, hex_color):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def bottom_border(p, color="2E3F5C", size=12):
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    b = OxmlElement("w:bottom")
    b.set(qn("w:val"), "single")
    b.set(qn("w:sz"), str(size))
    b.set(qn("w:space"), "1")
    b.set(qn("w:color"), color)
    bdr.append(b)
    pPr.append(bdr)


def first_cell_paragraph(cell):
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.15
    return p


class Sheet:
    """依 teacher 旗標輸出學生版（空白底線）或教師版（朱紅加粗答案）。"""

    def __init__(self, teacher: bool):
        self.teacher = teacher
        self.doc = Document()
        sec = self.doc.sections[0]
        sec.page_width, sec.page_height = Cm(21), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(1.3)
        sec.left_margin = sec.right_margin = Cm(1.6)
        normal = self.doc.styles["Normal"]
        normal.font.name = BODY_FONT
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        normal.font.size = Pt(BODY_PT)

    # 內嵌文字：字串＝一般文字；("A", 答案, 空白寬度)＝填空
    def rich(self, p, parts, size=BODY_PT):
        for part in parts:
            if isinstance(part, tuple):
                _, answer, width = part
                if self.teacher:
                    set_font(p.add_run(f" {answer} "), size, bold=True, color=RED, underline=True)
                else:
                    set_font(p.add_run("\u3000" * width), size, underline=True)
            else:
                set_font(p.add_run(part), size)
        return p

    def line(self, parts, size=BODY_PT, after=2, indent=0.0, before=0):
        p = para(self.doc, after=after, before=before)
        if indent:
            p.paragraph_format.left_indent = Cm(indent)
        return self.rich(p, parts if isinstance(parts, list) else [parts], size)

    def heading(self, text, sub=""):
        p = para(self.doc, after=3, before=5)
        set_font(p.add_run(text), 13.5, bold=True, color=NAVY, font=HEAD_FONT)
        if sub:
            set_font(p.add_run("　" + sub), 9.5, color=GRAY, font=HEAD_FONT)
        bottom_border(p)
        return p

    # 開放題：學生版畫書寫線，教師版顯示朱紅示範答案
    def answer_lines(self, n, example, indent=0.4):
        if self.teacher:
            p = para(self.doc, after=3)
            p.paragraph_format.left_indent = Cm(indent)
            set_font(p.add_run("【參考答案】"), 10.5, bold=True, color=RED)
            set_font(p.add_run(example), 10.5, bold=True, color=RED, underline=True)
            return
        for _ in range(n):
            p = para(self.doc, after=0, line=1.0)
            p.paragraph_format.left_indent = Cm(indent)
            set_font(p.add_run(" "), 15)
            bottom_border(p, color="999999", size=4)
        para(self.doc, after=1, line=0.6)

    def note(self, text):
        """教師版專屬的語譯／教學提示，學生版不輸出。"""
        if not self.teacher:
            return
        p = para(self.doc, after=2)
        p.paragraph_format.left_indent = Cm(0.4)
        set_font(p.add_run("【教學提示】" + text), 9.5, color=RED)

    def box(self, lines, fill="F6EFE1"):
        t = self.doc.add_table(rows=1, cols=1)
        t.style = "Table Grid"
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.cell(0, 0)
        shade(c, fill)
        for i, ln in enumerate(lines):
            p = first_cell_paragraph(c) if i == 0 else c.add_paragraph()
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.2
            if isinstance(ln, tuple):  # (粗體標題, 內文)
                set_font(p.add_run(ln[0]), 10.5, bold=True, color=NAVY)
                set_font(p.add_run(ln[1]), 10.5)
            else:
                set_font(p.add_run(ln), 10.5)
        para(self.doc, after=1, line=0.6)
        return t

    def page_break(self):
        self.doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(path)
        print("✔", path)


# ---------------------------------------------------------------- 內容
def build(teacher: bool):
    s = Sheet(teacher)
    doc = s.doc

    # ===== 標題區 =====
    head = doc.add_table(rows=1, cols=2)
    head.alignment = WD_TABLE_ALIGNMENT.CENTER
    left, right = head.cell(0, 0), head.cell(0, 1)
    left.width, right.width = Cm(14.6), Cm(3.2)
    p = first_cell_paragraph(left)
    set_font(p.add_run("AI 古人聊天室｜與孟子對話"), 19, bold=True, color=NAVY, font=HEAD_FONT)
    p = left.add_paragraph()
    tag = "教師解答版" if teacher else "自學學習單"
    set_font(p.add_run(f"〈亞聖的雄辯教室〉{tag}"), 12.5, bold=True, color=RED if teacher else GRAY, font=HEAD_FONT)
    p = left.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    set_font(p.add_run("班級：＿＿＿＿　座號：＿＿＿　姓名：＿＿＿＿＿＿＿"), 11)
    rp = first_cell_paragraph(right)
    rp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if AVATAR.exists():
        rp.add_run().add_picture(str(AVATAR), width=Cm(2.6))

    # ===== 一、作家履歷小檔案 =====
    s.heading("一、作家履歷小檔案", "點選聊天室「👦 自我介紹」並參考課本完成")
    rows = [
        ("姓名／字", ["姓孟，名", ("A", "軻", 2), "，相傳字子輿"]),
        ("時代／籍貫", [("A", "戰國", 3), "時代　", ("A", "鄒", 2), "國人（今山東鄒城）"]),
        ("尊　　稱", ["後世尊為「", ("A", "亞聖", 3), "」，與孔子並稱「孔孟」"]),
        ("師　　承", ["受業於", ("A", "子思", 3), "之門人，私淑", ("A", "孔子", 3)]),
        ("核心主張", ["人性本", ("A", "善", 2), "、行", ("A", "仁政", 3), "、民本思想、養浩然之氣"]),
        ("著　　作", ["《孟子》共", ("A", "七", 2), "篇，與《論語》《大學》《中庸》合稱「", ("A", "四書", 3), "」"]),
        ("文章風格", ["氣勢磅礴、善用", ("A", "比喻", 3), "與寓言、長於論辯"]),
    ]
    t = doc.add_table(rows=len(rows), cols=2)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (k, v) in enumerate(rows):
        c0, c1 = t.cell(i, 0), t.cell(i, 1)
        c0.width, c1.width = Cm(3.0), Cm(14.8)
        shade(c0, "E8EDF5")
        set_font(first_cell_paragraph(c0).add_run(k), 10.5, bold=True, color=NAVY)
        s.rich(first_cell_paragraph(c1), v, 10.5)

    # ===== 二、聊天室初探任務 =====
    s.heading("二、聊天室初探任務", "點「👦 自我介紹」「📜 得意作品」，邊聊邊記")
    s.line(["1. 孟母三遷：孟家依序搬到", ("A", "墓地", 3), "旁 →", ("A", "市集", 3), "旁 →", ("A", "學宮", 3), "旁，才定居下來。"])
    s.line(["2. 孟母「斷織」想告訴孟子什麼道理？"])
    s.answer_lines(1, "讀書半途而廢，就像把織到一半的布割斷，前功盡棄。")
    s.line(["3. 四端之心：惻隱之心是", ("A", "仁", 2), "之端；羞惡之心是", ("A", "義", 2),
            "之端；辭讓之心是", ("A", "禮", 2), "之端；是非之心是", ("A", "智", 2), "之端。"])
    s.line(["4. 孟子用「孺子將入於井」說明人人皆有", ("A", "惻隱", 3), "之心，所以人性本善。"])

    # ===== 三、核心課文精讀 =====
    s.heading("三、核心課文精讀", "〈魚我所欲也〉〈生於憂患，死於安樂〉")
    s.box([
        ("〈魚我所欲也〉　", "魚，我所欲也；熊掌，亦我所欲也。二者不可得兼，舍魚而取熊掌者也。"
         "生，亦我所欲也；義，亦我所欲也。二者不可得兼，舍生而取義者也。"),
        ("〈生於憂患〉　", "故天將降大任於是人也，必先苦其心志，勞其筋骨，餓其體膚，空乏其身，"
         "行拂亂其所為，所以動心忍性，曾益其所不能。……然後知生於憂患，而死於安樂也。"),
    ])
    s.note("語譯：魚和熊掌都是我想要的，不能兼得就捨魚取熊掌；生命和道義都是我想要的，不能兼得就捨生命而取道義。"
           "上天要把重任交給一個人，一定先使他內心痛苦、筋骨勞累、身體飢餓、生活貧困，做事總不順利，"
           "藉此激勵他的心志、堅韌他的性情，增加他原本沒有的能力。")
    s.line(["1. 孟子用「魚」比喻", ("A", "生命", 3), "，用「熊掌」比喻", ("A", "道義", 3),
            "；這種先說淺顯事物再引出主題的寫法稱為", ("A", "類比（譬喻）", 5), "。"])
    s.line(["2. 「曾益其所不能」的「曾」通「", ("A", "增", 2), "」，整句意思是：", ("A", "增加他原本不具備的能力", 10), "。"])
    s.line(["3. 生活連結：寫下一次你「生於憂患」——因困難而成長的經驗。"])
    s.answer_lines(2, "國二數學不及格後，每天放學多做十題並請教同學，期末進步了二十分，也學會不輕易放棄。（言之成理即可）")

    s.page_break()

    # ===== 四、八卦冷知識與典故考證 =====
    s.heading("四、八卦冷知識與典故考證", "點「🍵 八卦趣聞」「👥 諸子朋友圈」完成")
    idioms = [
        ("揠苗助長", "宋人嫌禾苗長太慢，一株株往上拔，結果「苗則槁矣」", "違反自然規律、急於求成，反而壞事"),
        ("五十步笑百步", "逃跑五十步的士兵嘲笑逃跑一百步的", "本質相同，只是程度不同"),
        ("緣木求魚", "孟子批評齊宣王想用武力稱霸天下", "方法錯誤，必定達不到目的"),
        ("王顧左右而言他", "孟子追問「四境之內不治，則如之何」", "無法回答時轉移話題"),
        ("以羊易牛", "齊宣王不忍見牛發抖，下令換羊祭鐘", "出於不忍之心；也指以小換大"),
    ]
    t = doc.add_table(rows=len(idioms) + 1, cols=3)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Cm(3.4), Cm(7.6), Cm(6.8)]
    for j, h in enumerate(["成語", "聊天室裡的故事", "比喻意思（請自行填寫）"]):
        c = t.cell(0, j)
        c.width = widths[j]
        shade(c, "2E3F5C")
        set_font(first_cell_paragraph(c).add_run(h), 10.5, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF))
    for i, (idiom, story, meaning) in enumerate(idioms, start=1):
        cells = [t.cell(i, j) for j in range(3)]
        for j, c in enumerate(cells):
            c.width = widths[j]
        set_font(first_cell_paragraph(cells[0]).add_run(idiom), 10.5, bold=True, color=NAVY)
        set_font(first_cell_paragraph(cells[1]).add_run(story), 10)
        p = first_cell_paragraph(cells[2])
        if teacher:
            set_font(p.add_run(meaning), 10, bold=True, color=RED, underline=True)
        else:
            set_font(p.add_run(" "), 10)

    s.line(["是非判斷（○／×）："], before=3, after=1)
    facts = [
        ("○", "孟子出生時，孔子已過世一百多年，所以孟子自稱「私淑」孔子。"),
        ("×", "孟子與告子辯論人性，告子主張「人性本善」。（告子主張性無善無不善）"),
        ("○", "淳于髡問「嫂溺援之以手乎」，孟子回答這是懂得變通的「權」。"),
        ("○", "明太祖讀到「君之視臣如土芥，則臣視君如寇讎」大怒，曾刪節《孟子》。"),
    ]
    for ans, f in facts:
        text = f if teacher else f.split("（")[0]
        s.line(["（", ("A", ans, 2), "）" + text], size=10.5, indent=0.3, after=1)

    # ===== 五、友情大考驗 =====
    s.heading("五、友情大考驗闖關", "先在聊天室按「🏆 友情考核」暖身，再作答")
    mcq = [
        ("C", "宋人「揠苗助長」的結果是：(A) 禾苗特別高 (B) 收成加倍 (C) 禾苗枯死 (D) 鄰人仿效"),
        ("B", "下列何者最能代表孟子的民本思想？(A) 王何必曰利 (B) 民為貴，社稷次之，君為輕 "
              "(C) 盡信書，則不如無書 (D) 我善養吾浩然之氣"),
        ("A", "孟子說「予豈好辯哉？予不得已也」，主要是為了：(A) 駁斥楊朱、墨翟之說，捍衛儒家 "
              "(B) 爭取齊王封官 (C) 與告子比賽口才 (D) 替孟母辯護"),
    ]
    for i, (ans, q) in enumerate(mcq, 1):
        s.line(["（", ("A", ans, 2), f"）{i}. {q}"], size=10.5, after=1)
    s.line(["典故填空："], before=3, after=1)
    s.line(["① 民為", ("A", "貴", 2), "，社稷次之，君為", ("A", "輕", 2), "。　",
            "② 天時不如", ("A", "地利", 3), "，地利不如", ("A", "人和", 3), "。"], size=10.5, indent=0.3, after=1)
    s.line(["③ 富貴不能", ("A", "淫", 2), "，貧賤不能", ("A", "移", 2), "，威武不能", ("A", "屈", 2),
            "，此之謂大丈夫。　④ 得道者", ("A", "多", 2), "助，失道者", ("A", "寡", 2), "助。"],
           size=10.5, indent=0.3, after=1)

    # ===== 六、人機協作任務 =====
    s.heading("六、人機協作任務：孟夫子辯論擂台通關記", "開啟「🤖 AI 對話」→ 按「⚔️ 辯論擂台」")
    s.line(["1. 我的主張（例：「讀書只是為了考試」「人性本惡」）："], size=10.5)
    s.answer_lines(1, "「人性本惡，小孩都會搶玩具。」")
    s.line(["2. 孟子用了什麼比喻或反問來回應我？"], size=10.5)
    s.answer_lines(2, "孟子說：「人性向善，就像水往低處流。小孩搶玩具是欲望遮蔽了本心，但看到別人哭也會不忍，這就是惻隱之心。」")
    s.line(["3. 辯論後，我同意／不同意孟子嗎？請說明理由（30 字以上）："], size=10.5)
    s.answer_lines(2, "部分同意。人確實有同情心，但若沒有教育與好環境，善端也可能被埋沒——這也是孟母三遷的道理。（觀點自由，理由充分即可）")

    name = "孟子學習單_教師解答版.docx" if teacher else "孟子學習單_學生版.docx"
    s.save(OUT / name)


if __name__ == "__main__":
    build(teacher=False)
    build(teacher=True)
